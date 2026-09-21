#!/usr/bin/env python3
"""SDOC hackathon prototype: classify emails, then compare SI and BL files.

For every category, the program reports:
  - which keywords matched the subject
  - which keywords matched the current email body
  - which strong phrases matched
  - the resulting score

Email classification uses transparent keyword scores.  Only emails classified
as BL_COMPARISON continue to attachment extraction, field-label matching, and
strict value comparison.  A detailed JSON report keeps every decision easy to
inspect.
"""

from __future__ import annotations

import argparse
import io
import importlib.util
import json
import re
import sys
import unicodedata
from collections import Counter
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
from pathlib import Path


DEFAULT_BUNDLE = Path("/Users/jonathanong/Documents/sdoc-hackathon-bundle")

# Grammar words and email greetings do not provide useful category evidence.
# Words such as "with" and "once" are deliberately ignored.
STOPWORDS = {
    "a", "about", "after", "all", "am", "an", "and", "any", "are", "as",
    "at", "be", "because", "been", "before", "being", "between", "both",
    "but", "by", "can", "could", "dear", "did", "do", "does", "doing",
    "down", "during", "each", "for", "from", "further", "good", "had",
    "has", "have", "having", "he", "hello", "her", "here", "hers", "hi",
    "him", "his", "how", "i", "if", "in", "into", "is", "it", "its",
    "itself", "just", "kindly", "me", "more", "most", "my", "no", "nor",
    "not", "of", "off", "on", "once", "only", "or", "other", "our", "out",
    "over", "own", "please", "same", "she", "should", "so", "some", "such",
    "than", "that", "the", "their", "them", "then", "there", "these", "they",
    "this", "those", "through", "to", "too", "under", "until", "up", "very",
    "was", "we", "were", "what", "when", "where", "which", "while", "who",
    "why", "will", "with", "would", "you", "your", "bye", "thanks", "thank",
}

# All keywords contribute one match. Subject matches receive more weight later.
# Keep the lists focused: a large list would unfairly give one category more
# opportunities to score than the others.
CATEGORY_KEYWORDS = {
    "SI_REQUEST": {
        "request", "si", "needed", "shipping", "instruction", "revert",
        "draft", "bl", "available", "send",
    },
    "BL_COMPARISON": {
        "confirm", "docs", "attached", "si", "draft", "bl", "check",
        "details", "verify", "matches", "amend", "compare", "aie", "afemy",
        "afptme", "afrt",
    },
    "INVOICE_QUERY": {
        "invoice", "billing", "charges", "missing", "gr", "query",
        "freight", "payment", "cost", "fee", "amount", "breakdown",
    },
    "GENERAL": {
        "update", "summary", "berthing", "report", "schedule", "reminder",
        "operations", "holiday", "loading", "completed", "vessel", "pending",
        "outstanding", "approval", "leave", "release", "daily",
    },
    "SPAM": {
        "won", "prize", "claim", "gift", "bitcoin", "guaranteed", "urgent",
        "account", "survey", "click", "winner", "investment", "parcel",
        "mailbox", "storage", "deactivation", "crypto", "weird", "trick",
        "iphone", "customs", "congratulations", "returns",
    },
}

# These combinations are stronger evidence than their individual words.
CATEGORY_PHRASES = {
    "SI_REQUEST": {
        "request si", "si needed", "shipping instruction",
        "revert with draft bl", "send the si",
    },
    "BL_COMPARISON": {
        "confirm docs", "attached are the si and draft bl", "check the details",
        "verify the bl", "matches the si", "amend bl", "request bl draft",
    },
    "INVOICE_QUERY": {
        "invoice query", "local charges", "missing gr", "freight charges",
        "invoice amount",
    },
    "GENERAL": {
        "update summary", "berthing report", "operations update",
        "holiday reminder", "pending bl release", "outstanding bl",
        "time off request", "daily berthing report",
    },
    "SPAM": {
        "you have won", "claim your", "gift card", "bitcoin investment",
        "verify account", "verify your account", "guaranteed returns",
        "weird trick", "storage limit", "free iphone", "unpaid customs",
        "monthly draw", "click here",
    },
}

SUBJECT_KEYWORD_WEIGHT = 2
BODY_KEYWORD_WEIGHT = 1
SUBJECT_PHRASE_BONUS = 2
BODY_PHRASE_BONUS = 1


# Canonical document fields and the labels that the supplied documents use.
# The prototype never compares a value until its label has been assigned with
# enough confidence.
FIELD_ALIASES = {
    "shipper": [
        "Shipper", "Shipper/Exporter", "Shipper (Principal or Seller)",
        "SHIPPER",
    ],
    "consignee": [
        "Consignee", "Consignee (Non-Negotiable)", "CONSIGNEE",
        "To the Order of",
    ],
    "notify_party": [
        "Notify Party", "Notify", "Notify Party/Intermediate Consignee",
        "NOTIFY PARTY",
    ],
    "port_of_loading": [
        "Port of Loading", "Port of Loading (POL)", "Load Port", "POL",
        "PORT OF LOADING",
    ],
    "port_of_discharge": [
        "Port of Discharge", "Port of Discharge (POD)", "Discharge Port",
        "POD", "PORT OF DISCHARGE",
    ],
    "container_count": [
        "No. of Containers", "Total Containers",
        "No. of Containers or Packages", "Container Count",
    ],
    "gross_weight_kg": [
        "Gross Weight (KG)", "Gross Wt (kgs)", "Gross Weight毛重(KGS)",
        "GROSS WEIGHT",
    ],
}

MIN_FIELD_SCORE = 30.0
MIN_FIELD_MARGIN = 10.0
LABEL_CONNECTORS = {"a", "and", "of", "or", "the"}
WRONG_DOCUMENT_MARKERS = (
    "commercial invoice",
    "packing list",
    "certificate of origin",
)
MISSING_VALUE_WORDS = {"", "tba", "tbc", "n/a", "na", "nil", "none"}


def ascii_words(text: object) -> list[str]:
    """Return lowercase ASCII words while tolerating punctuation and Unicode."""

    normalised = unicodedata.normalize("NFKC", str(text or "")).casefold()
    return re.findall(r"[a-z0-9]+", normalised)


def label_words(text: object) -> list[str]:
    """Normalise common label abbreviations without changing their meaning."""

    replacements = {
        "containers": "container",
        "kgs": "kg",
        "no": "number",
        "packages": "package",
        "wt": "weight",
    }
    return [replacements.get(word, word) for word in ascii_words(text)]


def meaningful_label_words(text: object) -> list[str]:
    return [word for word in label_words(text) if word not in LABEL_CONNECTORS]


def label_acronyms(text: object) -> set[str]:
    """Return acronym forms, including POD and PD for Port of Discharge."""

    words = label_words(text)
    if len(words) == 1 and words[0].isalpha() and len(words[0]) <= 5:
        return {words[0]}
    with_connectors = "".join(word[0] for word in words if word)
    without_connectors = "".join(
        word[0] for word in words if word and word not in LABEL_CONNECTORS
    )
    return {value for value in (with_connectors, without_connectors) if value}


def score_label_against_alias(label: object, alias: str) -> dict:
    """Score one extracted label against one known alias.

    Exact aliases are strongest.  Otherwise, whole words and acronyms lead the
    score, matching initials support it, character similarity gives a small
    bonus, and different words/initials receive small penalties.
    """

    left_words = label_words(label)
    right_words = label_words(alias)
    if left_words and left_words == right_words:
        return {
            "score": 100.0,
            "exact_alias": True,
            "whole_word_points": 0.0,
            "acronym_points": 0.0,
            "initial_points": 0.0,
            "character_points": 0.0,
            "different_word_penalty": 0.0,
        }

    left_meaningful = set(meaningful_label_words(label))
    right_meaningful = set(meaningful_label_words(alias))
    shared_words = left_meaningful & right_meaningful
    different_words = left_meaningful ^ right_meaningful

    whole_word_points = 15.0 * len(shared_words)
    different_word_penalty = -5.0 * len(different_words)

    left_acronyms = label_acronyms(label)
    right_acronyms = label_acronyms(alias)
    shared_acronyms = left_acronyms & right_acronyms
    acronym_points = (
        30.0
        if any(len(acronym) >= 2 for acronym in shared_acronyms)
        else 0.0
    )

    # An acronym such as PD is not a genuinely different word from the full
    # label Port of Discharge, so do not penalise it as one when they agree.
    if acronym_points:
        different_word_penalty = 0.0

    initial_scores = []
    for left_acronym in left_acronyms:
        for right_acronym in right_acronyms:
            score = 0.0
            for left_initial, right_initial in zip(left_acronym, right_acronym):
                score += 5.0 if left_initial == right_initial else -5.0
            initial_scores.append(score)
    initial_points = max(initial_scores, default=0.0)

    left_text = "".join(left_words)
    right_text = "".join(right_words)
    character_points = (
        round(SequenceMatcher(None, left_text, right_text).ratio() * 10.0, 2)
        if left_text and right_text
        else 0.0
    )
    score = max(
        0.0,
        whole_word_points
        + acronym_points
        + initial_points
        + character_points
        + different_word_penalty,
    )
    return {
        "score": round(score, 2),
        "exact_alias": False,
        "whole_word_points": whole_word_points,
        "acronym_points": acronym_points,
        "initial_points": initial_points,
        "character_points": character_points,
        "different_word_penalty": different_word_penalty,
    }


def classify_field_label(label: object) -> dict:
    """Assign a label only when the best field is strong and unambiguous."""

    candidates = []
    for field, aliases in FIELD_ALIASES.items():
        alias_results = [
            (alias, score_label_against_alias(label, alias)) for alias in aliases
        ]
        best_alias, breakdown = max(
            alias_results, key=lambda item: item[1]["score"]
        )
        candidates.append(
            {
                "field": field,
                "best_alias": best_alias,
                **breakdown,
            }
        )

    candidates.sort(key=lambda item: item["score"], reverse=True)
    best = candidates[0]
    second = candidates[1]
    margin = round(best["score"] - second["score"], 2)
    accepted = best["score"] >= MIN_FIELD_SCORE and margin >= MIN_FIELD_MARGIN
    return {
        "field": best["field"] if accepted else None,
        "accepted": accepted,
        "best_score": best["score"],
        "second_score": second["score"],
        "margin": margin,
        "best_alias": best["best_alias"],
        "candidates": candidates,
    }


def remove_email_boilerplate(body: str) -> str:
    """Keep the current message and remove signatures/forwarded history."""

    marker = re.search(
        r"(?im)^\s*(?:best regards|kind regards|regards|thanks and regards|best),?\s*$"
        r"|^_{5,}\s*$"
        r"|^-{5,}\s*original message\s*-{5,}$",
        body,
    )
    return body[: marker.start()] if marker else body


def normalise_text(text: str) -> str:
    """Lowercase text and turn punctuation/underscores into spaces."""

    return " ".join(re.findall(r"[a-z0-9]+", text.lower()))


def extract_words(text: str) -> list[str]:
    """Break text into useful words while keeping short terms such as SI/BL."""

    words = re.findall(r"[a-z]+", text.lower())
    return [
        word
        for word in words
        if word not in STOPWORDS and (len(word) >= 3 or word in {"si", "bl", "gr"})
    ]


def count_matches(words: list[str], category: str) -> dict[str, int]:
    """Return the occurrence count for every matching category keyword."""

    occurrences = Counter(words)
    return {
        keyword: occurrences[keyword]
        for keyword in sorted(CATEGORY_KEYWORDS[category])
        if occurrences[keyword] > 0
    }


def find_phrases(text: str, category: str) -> list[str]:
    normalised = normalise_text(text)
    return sorted(
        phrase
        for phrase in CATEGORY_PHRASES[category]
        if phrase in normalised
    )


def score_category(subject: str, body: str, category: str) -> dict:
    subject_matches = count_matches(extract_words(subject), category)
    body_matches = count_matches(extract_words(body), category)
    subject_phrases = find_phrases(subject, category)
    body_phrases = find_phrases(body, category)

    # Repeated words are shown in the dictionaries above, but classification
    # uses the number of UNIQUE matched keywords so repetition cannot dominate.
    subject_keyword_score = len(subject_matches) * SUBJECT_KEYWORD_WEIGHT
    body_keyword_score = len(body_matches) * BODY_KEYWORD_WEIGHT
    phrase_bonus = (
        len(subject_phrases) * SUBJECT_PHRASE_BONUS
        + len(body_phrases) * BODY_PHRASE_BONUS
    )

    return {
        "subject_matches": subject_matches,
        "body_matches": body_matches,
        "subject_match_count": len(subject_matches),
        "body_match_count": len(body_matches),
        "subject_phrases": subject_phrases,
        "body_phrases": body_phrases,
        "subject_keyword_score": subject_keyword_score,
        "body_keyword_score": body_keyword_score,
        "phrase_bonus": phrase_bonus,
        "total_score": subject_keyword_score + body_keyword_score + phrase_bonus,
    }


def classify_email(email: dict) -> dict:
    subject = email.get("subject", "")
    body = remove_email_boilerplate(email.get("body", ""))

    scores = {
        category: score_category(subject, body, category)
        for category in CATEGORY_KEYWORDS
    }

    ranked = sorted(
        scores,
        key=lambda category: (
            scores[category]["total_score"],
            scores[category]["subject_keyword_score"],
            scores[category]["phrase_bonus"],
        ),
        reverse=True,
    )

    best = ranked[0]
    second = ranked[1]
    best_score = scores[best]["total_score"]
    second_score = scores[second]["total_score"]

    # With no evidence, GENERAL is a safer baseline than choosing randomly.
    if best_score == 0:
        best = "GENERAL"

    keyword_category = best
    routing_reason = "keyword_score"
    attachments = email.get("attachments", [])
    has_si = any(re.search(r"_SI\.[^.]+$", path, re.I) for path in attachments)
    has_bl = any(re.search(r"_BL\.[^.]+$", path, re.I) for path in attachments)
    if has_si and has_bl:
        # A named SI+BL pair is stronger evidence of a comparison task than a
        # subject such as "request BL draft", which can describe either stage.
        best = "BL_COMPARISON"
        routing_reason = "si_bl_attachment_pair"

    best_score = scores[best]["total_score"]
    second_score = max(
        scores[category]["total_score"]
        for category in scores
        if category != best
    )

    return {
        "category": best,
        "keyword_category": keyword_category,
        "routing_reason": routing_reason,
        "winning_score": best_score,
        "score_margin": best_score - second_score,
        "ambiguous": (
            routing_reason == "keyword_score"
            and best_score > 0
            and best_score == second_score
        ),
        "scores": scores,
    }


def attachment_suffix(path: str) -> str:
    return Path(path).suffix.casefold()


def read_attachment(inbox, path: str) -> dict:
    """Extract text and label/value pairs from TXT, PDF, DOCX, or XLSX."""

    result = {
        "path": path,
        "format": attachment_suffix(path).lstrip("."),
        "text": "",
        "lines": [],
        "pairs": [],
        "error": None,
    }
    try:
        data = inbox.read_bytes(path)
        suffix = attachment_suffix(path)

        if suffix == ".txt":
            text = data.decode("utf-8", errors="replace")
            result["text"] = text
            result["lines"] = text.splitlines()

        elif suffix == ".pdf":
            try:
                import pdfplumber
            except ImportError as error:
                raise RuntimeError(
                    "PDF support needs pdfplumber (pip install pdfplumber)."
                ) from error
            pages = []
            pairs = []
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                for page in pdf.pages:
                    pages.append(page.extract_text() or "")
                    # In the supplied PDF layout, labels are bold and values
                    # are regular text.  Separating by font avoids merged text
                    # when a long label overlaps the value column visually.
                    rows = {}
                    for char in page.chars:
                        rows.setdefault(round(float(char["top"]), 1), []).append(char)
                    for row_chars in rows.values():
                        bold = sorted(
                            (
                                char
                                for char in row_chars
                                if "bold" in str(char.get("fontname", "")).casefold()
                            ),
                            key=lambda char: char["x0"],
                        )
                        regular = sorted(
                            (
                                char
                                for char in row_chars
                                if "bold" not in str(char.get("fontname", "")).casefold()
                            ),
                            key=lambda char: char["x0"],
                        )
                        label = "".join(char["text"] for char in bold).strip()
                        value = "".join(char["text"] for char in regular).strip()
                        if label and value:
                            pairs.append((label, value))
            text = "\n".join(pages)
            result["text"] = text
            result["lines"] = text.splitlines()
            result["pairs"] = pairs

        elif suffix == ".docx":
            try:
                from docx import Document
            except ImportError as error:
                raise RuntimeError(
                    "DOCX support needs python-docx (pip install python-docx)."
                ) from error
            document = Document(io.BytesIO(data))
            lines = [p.text for p in document.paragraphs if p.text.strip()]
            pairs = []
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    lines.append(" | ".join(cells))
                    if len(cells) >= 2:
                        pairs.append((cells[0], cells[1]))
            result["lines"] = lines
            result["pairs"] = pairs
            result["text"] = "\n".join(lines)

        elif suffix == ".xlsx":
            try:
                from openpyxl import load_workbook
            except ImportError as error:
                raise RuntimeError(
                    "XLSX support needs openpyxl (pip install openpyxl)."
                ) from error
            workbook = load_workbook(
                io.BytesIO(data), read_only=True, data_only=True
            )
            lines = []
            pairs = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value).strip() for value in row]
                    if not any(values):
                        continue
                    lines.append(" | ".join(values))
                    if len(values) >= 2:
                        pairs.append((values[0], values[1]))
            workbook.close()
            result["lines"] = lines
            result["pairs"] = pairs
            result["text"] = "\n".join(lines)

        else:
            raise ValueError(f"Unsupported attachment format: {suffix or '(none)'}")

        if not result["text"].strip() and not result["pairs"]:
            raise ValueError("The attachment contains no readable text.")
    except Exception as error:  # one bad attachment must not stop the inbox
        result["error"] = f"{type(error).__name__}: {error}"
    return result


def label_prefix_pattern(alias: str) -> re.Pattern:
    """Build a tolerant start-of-line pattern from an alias's ASCII words."""

    words = ascii_words(alias)
    joined = r"[^A-Za-z0-9]+".join(re.escape(word) for word in words)
    return re.compile(
        r"^\s*" + joined + r"(?:[^A-Za-z0-9]+|$)",
        flags=re.IGNORECASE,
    )


ALIAS_PREFIXES = sorted(
    (
        (alias, field, label_prefix_pattern(alias))
        for field, aliases in FIELD_ALIASES.items()
        for alias in aliases
    ),
    key=lambda item: len(ascii_words(item[0])),
    reverse=True,
)


def store_field_candidate(
    fields: dict, label: object, value: object, source: str
) -> None:
    """Classify and store one label/value pair, preferring usable values."""

    decision = classify_field_label(label)
    field = decision["field"]
    if field is None:
        return
    candidate = {
        "raw_label": str(label or "").strip(),
        "raw_value": str(value or "").strip(),
        "source": source,
        "label_match": decision,
    }
    previous = fields.get(field)
    if previous is None:
        fields[field] = candidate
        return
    previous_empty = is_missing_value(previous["raw_value"])
    candidate_empty = is_missing_value(candidate["raw_value"])
    if (previous_empty and not candidate_empty) or (
        previous_empty == candidate_empty
        and decision["best_score"] > previous["label_match"]["best_score"]
    ):
        fields[field] = candidate


def extract_document_fields(document: dict) -> dict:
    """Extract the seven required canonical fields from one document."""

    fields = {}
    uncertain_labels = []

    # Tables and spreadsheets give the label and value in adjacent cells.
    for label, value in document["pairs"]:
        decision = classify_field_label(label)
        if decision["accepted"]:
            store_field_candidate(fields, label, value, "table")
        elif str(label).strip():
            uncertain_labels.append(
                {"label": str(label), "source": "table", "decision": decision}
            )

    # TXT often uses "Label: value".  Extracted PDF text often uses
    # "Label value", so known aliases are also detected at the line start.
    for line_number, original_line in enumerate(document["lines"], start=1):
        line = original_line.strip()
        if not line:
            continue
        # Preserve "Total Containers" because it is a real alias.  A second
        # version without TOTAL handles PDF lines such as "TOTAL Gross Wt".
        line_variants = [line]
        without_total = re.sub(r"(?i)^total\s+", "", line).strip()
        if without_total != line:
            line_variants.append(without_total)

        stored = False
        for candidate_line in line_variants:
            if ":" in candidate_line:
                label, value = candidate_line.split(":", 1)
                decision = classify_field_label(label)
                if decision["accepted"]:
                    store_field_candidate(fields, label, value, f"line {line_number}")
                    stored = True
                    break

            for alias, _field, pattern in ALIAS_PREFIXES:
                match = pattern.match(candidate_line)
                if not match:
                    continue
                value = candidate_line[match.end() :].lstrip(" :-|\t")
                store_field_candidate(fields, alias, value, f"line {line_number}")
                stored = True
                break
            if stored:
                break

    return {
        "fields": fields,
        "missing_fields": [field for field in FIELD_ALIASES if field not in fields],
        "uncertain_labels": uncertain_labels,
    }


def is_missing_value(value: object) -> bool:
    """Recognise blanks and explicit placeholders as missing values."""

    text = unicodedata.normalize("NFKC", str(value or "")).strip().casefold()
    compact = re.sub(r"\s+", "", text)
    if compact in MISSING_VALUE_WORDS:
        return True
    return bool(re.fullmatch(r"[_?\-.]*(?:kg|kgs|mt|mts)?", compact))


def normalise_text_value(value: object) -> str | None:
    if is_missing_value(value):
        return None
    text = unicodedata.normalize("NFKC", str(value)).casefold().replace("&", "and")
    pieces = re.findall(r"[a-z0-9]+", text)
    return "".join(pieces) or None


def normalise_container_count(value: object) -> int | None:
    if is_missing_value(value):
        return None
    text = unicodedata.normalize("NFKC", str(value)).casefold()
    before_x = re.search(r"\b(\d+)\s*[x×]\b", text)
    match = before_x or re.search(r"\b\d+\b", text)
    return int(match.group(1) if before_x else match.group(0)) if match else None


def normalise_weight_kg(value: object) -> Decimal | None:
    if is_missing_value(value):
        return None
    text = unicodedata.normalize("NFKC", str(value)).casefold()
    match = re.search(r"(?<![a-z0-9])\d[\d, ]*(?:\.\d+)?", text)
    if not match:
        return None
    number_text = re.sub(r"[, ]", "", match.group(0))
    try:
        number = Decimal(number_text)
    except InvalidOperation:
        return None
    if re.search(r"\b(?:mt|mts|ton|tons|tonne|tonnes)\b", text):
        number *= Decimal("1000")
    return number.normalize()


def normalise_field_value(field: str, value: object):
    if field == "container_count":
        return normalise_container_count(value)
    if field == "gross_weight_kg":
        return normalise_weight_kg(value)
    return normalise_text_value(value)


def json_value(value: object):
    """Convert normalised values (including Decimal) into JSON-safe evidence."""

    if isinstance(value, Decimal):
        return format(value, "f")
    return value


def compare_document_fields(si_extraction: dict, bl_extraction: dict) -> dict:
    """Compare every required field after field-specific normalisation."""

    comparisons = {}
    mismatches = []
    reviews = []
    for field in FIELD_ALIASES:
        si_item = si_extraction["fields"].get(field)
        bl_item = bl_extraction["fields"].get(field)
        si_raw = si_item["raw_value"] if si_item else None
        bl_raw = bl_item["raw_value"] if bl_item else None
        si_normalised = normalise_field_value(field, si_raw)
        bl_normalised = normalise_field_value(field, bl_raw)

        if si_normalised is None or bl_normalised is None:
            state = "REVIEW_NEEDED"
            reviews.append(field)
        elif si_normalised == bl_normalised:
            state = "MATCHED"
        else:
            state = "NOT_MATCHED"
            mismatches.append(field)

        comparisons[field] = {
            "result": state,
            "si_raw": si_raw,
            "bl_raw": bl_raw,
            "si_normalised": json_value(si_normalised),
            "bl_normalised": json_value(bl_normalised),
            "si_label": si_item["raw_label"] if si_item else None,
            "bl_label": bl_item["raw_label"] if bl_item else None,
        }

    # Review takes precedence: the document pair cannot be decided safely.
    if reviews:
        return {
            "status": "NEEDS_REVIEW",
            "review_reason": "missing_value",
            "has_defect": False,
            "defect_fields": [],
            "review_fields": reviews,
            "known_mismatches": mismatches,
            "comparisons": comparisons,
        }
    if mismatches:
        return {
            "status": "MISMATCH",
            "review_reason": None,
            "has_defect": True,
            "defect_fields": mismatches,
            "review_fields": [],
            "known_mismatches": mismatches,
            "comparisons": comparisons,
        }
    return {
        "status": "OK",
        "review_reason": None,
        "has_defect": False,
        "defect_fields": [],
        "review_fields": [],
        "known_mismatches": [],
        "comparisons": comparisons,
    }


def needs_attachment_review(email: dict) -> bool:
    """Distinguish a comparison missing files from a normal draft-BL request."""

    attachments = email.get("attachments", [])
    if len(attachments) == 1:
        return True
    if attachments:
        return False
    text = normalise_text(f"{email.get('subject', '')} {email.get('body', '')}")
    explicit_phrases = (
        "compare the si and draft bl",
        "attached si and draft bl",
        "attachments dropped",
        "attachment dropped",
        "draft bl is still missing",
        "missing attachment",
    )
    return any(phrase in text for phrase in explicit_phrases)


def find_si_bl_paths(attachments: list[str]) -> tuple[str | None, str | None]:
    si_path = next(
        (path for path in attachments if re.search(r"_SI\.[^.]+$", path, re.I)),
        None,
    )
    bl_path = next(
        (path for path in attachments if re.search(r"_BL\.[^.]+$", path, re.I)),
        None,
    )
    return si_path, bl_path


def review_decision(reason: str) -> dict:
    return {
        "status": "NEEDS_REVIEW",
        "review_reason": reason,
        "has_defect": False,
        "defect_fields": [],
    }


def analyse_bl_comparison(inbox, email: dict) -> tuple[dict, dict]:
    """Read and compare SI/BL attachments for one routed email."""

    attachments = email.get("attachments", [])
    details = {"attachments": attachments}
    if len(attachments) < 2:
        if needs_attachment_review(email):
            return review_decision("missing_attachment"), details
        return {
            "status": "OK",
            "review_reason": None,
            "has_defect": False,
            "defect_fields": [],
        }, details

    si_path, bl_path = find_si_bl_paths(attachments)
    if not si_path or not bl_path:
        return review_decision("missing_attachment"), details

    si_document = read_attachment(inbox, si_path)
    bl_document = read_attachment(inbox, bl_path)
    details["si_document"] = si_document
    details["bl_document"] = bl_document
    if si_document["error"] or bl_document["error"]:
        return review_decision("unreadable"), details

    combined_text = f"{si_document['text']}\n{bl_document['text']}".casefold()
    if any(marker in combined_text for marker in WRONG_DOCUMENT_MARKERS):
        return review_decision("wrong_doc_type"), details

    si_extraction = extract_document_fields(si_document)
    bl_extraction = extract_document_fields(bl_document)
    details["si_extraction"] = si_extraction
    details["bl_extraction"] = bl_extraction
    comparison = compare_document_fields(si_extraction, bl_extraction)
    details["comparison"] = comparison
    return {
        key: comparison[key]
        for key in ("status", "review_reason", "has_defect", "defect_fields")
    }, details


def process_email(inbox, email: dict) -> tuple[dict, dict]:
    classification = classify_email(email)
    category = classification["category"]
    document_details = None
    if category == "BL_COMPARISON":
        document_decision, document_details = analyse_bl_comparison(inbox, email)
    else:
        document_decision = {
            "status": "OK",
            "review_reason": None,
            "has_defect": False,
            "defect_fields": [],
        }

    submission = {"category": category, **document_decision}
    details = {
        "classification": classification,
        "final_decision": submission,
        "document_analysis": document_details,
    }
    return submission, details


def load_inbox(bundle: Path):
    """Load the provided Inbox class without copying the dataset."""

    sys.path.insert(0, str(bundle))
    from loader import Inbox

    return Inbox(str(bundle))


def ensure_attachment_dependencies(emails: list[dict]) -> None:
    """Fail clearly instead of silently treating missing libraries as bad files."""

    format_modules = {
        ".pdf": ("pdfplumber", "pdfplumber"),
        ".docx": ("docx", "python-docx"),
        ".xlsx": ("openpyxl", "openpyxl"),
    }
    used_suffixes = {
        attachment_suffix(path)
        for email in emails
        for path in email.get("attachments", [])
    }
    missing_packages = [
        package
        for suffix, (module, package) in format_modules.items()
        if suffix in used_suffixes and importlib.util.find_spec(module) is None
    ]
    if missing_packages:
        packages = " ".join(missing_packages)
        raise SystemExit(
            "Missing attachment-reading packages. Install them with:\n"
            f"  python3 -m pip install {packages}\n"
            "Then run this prototype again."
        )


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Classify SDOC emails, then compare SI and BL attachment fields for "
            "BL_COMPARISON emails."
        )
    )
    parser.add_argument("--bundle", type=Path, default=DEFAULT_BUNDLE)
    parser.add_argument(
        "--submission",
        type=Path,
        default=Path("prototype_submission.json"),
        help="compact result in the hackathon submission format",
    )
    parser.add_argument(
        "--report",
        type=Path,
        default=Path("prototype_detailed_report.json"),
        help="full scoring, extraction, and comparison evidence",
    )
    parser.add_argument(
        "--show",
        type=int,
        default=10,
        help="number of example decisions to print",
    )
    args = parser.parse_args()

    inbox = load_inbox(args.bundle)
    emails = list(inbox)
    ensure_attachment_dependencies(emails)
    submission = {}
    report = {}
    category_totals = Counter()
    status_totals = Counter()
    ambiguous_count = 0

    for email in emails:
        decision, details = process_email(inbox, email)
        email_id = email["email_id"]
        submission[email_id] = decision
        report[email_id] = details
        category_totals[decision["category"]] += 1
        status_totals[decision["status"]] += 1
        ambiguous_count += int(details["classification"]["ambiguous"])

    args.submission.write_text(json.dumps(submission, indent=2) + "\n")
    args.report.write_text(json.dumps(report, indent=2) + "\n")

    print(f"Processed {len(submission)} emails")
    print(f"Saved hackathon submission to {args.submission.resolve()}")
    print(f"Saved detailed evidence to {args.report.resolve()}")
    print("\nPredicted category totals:")
    for category in CATEGORY_KEYWORDS:
        print(f"  {category:<16} {category_totals[category]:>4}")
    print(f"  {'AMBIGUOUS':<16} {ambiguous_count:>4}")

    print("\nFinal status totals:")
    for status in ("OK", "MISMATCH", "NEEDS_REVIEW"):
        print(f"  {status:<16} {status_totals[status]:>4}")

    print("\nExample decisions:")
    for email_id in list(submission)[: max(args.show, 0)]:
        decision = submission[email_id]
        classification = report[email_id]["classification"]
        print(
            f"  {email_id}: {decision['category']} / {decision['status']} "
            f"(email score={classification['winning_score']}, "
            f"margin={classification['score_margin']})"
        )
        if decision["defect_fields"]:
            print(f"    different fields: {decision['defect_fields']}")
        if decision["review_reason"]:
            print(f"    review reason: {decision['review_reason']}")


if __name__ == "__main__":
    main()
